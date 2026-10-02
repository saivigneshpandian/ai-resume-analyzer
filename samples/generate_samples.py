"""Generate fictional sample resumes for testing the analyzer.

Usage: pip install reportlab python-docx && python3 generate_samples.py
"""
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer

# Each resume is a list of (kind, content) blocks: h1/p/h2/bullets.
WEAK_RESUME = [
    ("h1", "Priya Sharma"),
    ("p", "priya.sharma@example.com | +91 98765 43210 | github.com/priyasharma | Chennai, India"),
    ("h2", "Summary"),
    ("p", "Hard-working and passionate final-year Computer Science student and team player looking for opportunities to grow."),
    ("h2", "Education"),
    ("p", "B.E. Computer Science and Engineering, Anna University, 2022-2026, CGPA 8.4/10"),
    ("h2", "Experience"),
    ("p", "Software Development Intern, TechNova Solutions, May 2025 - July 2025"),
    ("bullets", [
        "Responsible for managing the team's tasks.",
        "Worked on backend development using Node.js and Express.",
        "Helped with testing the application.",
        "Building dashboards in React for internal users.",
    ]),
    ("h2", "Projects"),
    ("p", "Campus Event Portal - React, Node.js, MongoDB"),
    ("bullets", [
        "Made a website for college events where students can register.",
        "Used JWT for login.",
    ]),
    ("p", "Expense Tracker App - Python, Flask, SQLite"),
    ("bullets", ["Developed an app to track expenses with charts."]),
    ("h2", "Skills"),
    ("p", "JavaScript, Python, React, Node.js, Express, MongoDB, SQL, Git, HTML, CSS"),
    ("h2", "Achievements"),
    ("bullets", ["Participated in Smart India Hackathon 2024."]),
]

SPARSE_RESUME = [
    ("h1", "Arun Kumar"),
    ("p", "arun.k@example.com"),
    ("h2", "Education"),
    ("p", "B.Tech IT, 2026"),
    ("h2", "Skills"),
    ("p", "Java, C, HTML"),
]


def _long_resume():
    blocks = [
        ("h1", "Meera Iyer"),
        ("p", "meera.iyer@example.com | linkedin.com/in/meeraiyer | Bengaluru, India"),
        ("h2", "Summary"),
        ("p", "Backend engineer with 6 years of experience building payment and logistics platforms in Java, Go and Python."),
        ("h2", "Experience"),
    ]
    roles = [
        ("Senior Software Engineer, PayGrid", "2023 - Present"),
        ("Software Engineer II, ShipFast Logistics", "2021 - 2023"),
        ("Software Engineer, ShipFast Logistics", "2020 - 2021"),
        ("Associate Engineer, DataCraft Analytics", "2019 - 2020"),
    ]
    bullets = [
        "Designed and shipped a reconciliation service in Go processing 2M transactions/day with 99.98% accuracy.",
        "Led migration of 14 services from a monolith to Kubernetes, cutting deployment time from 2 hours to 12 minutes.",
        "Worked on improving the performance of database queries.",
        "Mentored 4 junior engineers through onboarding and code reviews.",
        "Was involved in on-call rotation and incident handling.",
        "Built a rate-limiting library adopted by 9 internal teams, reducing 429 incidents by 60%.",
        "Participated in sprint planning and daily standups.",
        "Introduced contract testing with Pact, catching 30+ breaking API changes before release.",
    ]
    for title, dates in roles:
        blocks.append(("p", f"{title}, {dates}"))
        blocks.append(("bullets", bullets))
    blocks += [
        ("h2", "Projects"),
        ("p", "Open-source: go-retry - a retry/backoff library with 1.2k GitHub stars."),
        ("p", "Personal: Home automation dashboard using Raspberry Pi, MQTT and Grafana."),
        ("h2", "Education"),
        ("p", "B.Tech Computer Science, NIT Trichy, 2015-2019"),
        ("h2", "Skills"),
        ("p", "Go, Java, Python, PostgreSQL, Redis, Kafka, Kubernetes, Docker, AWS, Terraform, gRPC, REST"),
        ("h2", "Certifications"),
        ("bullets", ["AWS Certified Solutions Architect - Associate", "Certified Kubernetes Application Developer"]),
    ]
    return blocks


def write_pdf(blocks, path):
    styles = getSampleStyleSheet()
    story = []
    for kind, content in blocks:
        if kind == "h1":
            story.append(Paragraph(content, styles["Title"]))
        elif kind == "h2":
            story.append(Spacer(1, 6))
            story.append(Paragraph(content, styles["Heading2"]))
        elif kind == "p":
            story.append(Paragraph(content, styles["Normal"]))
        elif kind == "bullets":
            story.append(ListFlowable([Paragraph(b, styles["Normal"]) for b in content], bulletType="bullet"))
    SimpleDocTemplate(path, pagesize=A4).build(story)


def write_docx(blocks, path):
    doc = Document()
    for kind, content in blocks:
        if kind == "h1":
            doc.add_heading(content, level=0)
        elif kind == "h2":
            doc.add_heading(content, level=1)
        elif kind == "p":
            doc.add_paragraph(content)
        elif kind == "bullets":
            for b in content:
                doc.add_paragraph(b, style="List Bullet")
    doc.save(path)


if __name__ == "__main__":
    write_pdf(WEAK_RESUME, "sample-resume.pdf")
    write_docx(WEAK_RESUME, "sample-resume.docx")
    write_pdf(SPARSE_RESUME, "sparse-resume.pdf")
    write_pdf(_long_resume(), "long-resume.pdf")
    print("Generated sample resumes.")
